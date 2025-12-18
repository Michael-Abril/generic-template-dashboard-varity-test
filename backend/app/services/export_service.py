"""
Export Service - Data Export Functionality for Dashboard
Supports CSV, PDF, JSON, and Excel exports
"""
import csv
import io
import json
import logging
from datetime import datetime
from typing import Dict, Any, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .pdf_service import pdf_service

logger = logging.getLogger(__name__)


class ExportService:
    """Service for exporting dashboard data in various formats"""

    def __init__(self):
        """Initialize export service"""
        logger.info("Export Service initialized")

    async def export_to_csv(
        self, data: List[Dict[str, Any]], filename: str = "export.csv"
    ) -> bytes:
        """
        Export data to CSV format

        Args:
            data: List of dictionaries to export
            filename: Output filename (for content-disposition header)

        Returns:
            CSV file as bytes
        """
        if not data:
            raise ValueError("No data to export")

        # Create CSV in memory
        output = io.StringIO()

        # Get headers from first row
        headers = list(data[0].keys())

        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()

        # Write all rows
        for row in data:
            # Convert datetime objects to strings
            cleaned_row = {}
            for key, value in row.items():
                if isinstance(value, datetime):
                    cleaned_row[key] = value.isoformat()
                elif isinstance(value, dict) or isinstance(value, list):
                    cleaned_row[key] = json.dumps(value)
                else:
                    cleaned_row[key] = value
            writer.writerow(cleaned_row)

        # Convert to bytes
        csv_bytes = output.getvalue().encode("utf-8")
        logger.info(f"Exported {len(data)} rows to CSV ({len(csv_bytes)} bytes)")

        return csv_bytes

    async def export_to_json(
        self, data: Dict[str, Any], filename: str = "export.json", pretty: bool = True
    ) -> bytes:
        """
        Export data to JSON format

        Args:
            data: Dictionary or list to export
            filename: Output filename
            pretty: Whether to pretty-print JSON

        Returns:
            JSON file as bytes
        """
        if not data:
            raise ValueError("No data to export")

        # Convert datetime objects to strings recursively
        def clean_data(obj):
            if isinstance(obj, datetime):
                return obj.isoformat()
            elif isinstance(obj, dict):
                return {k: clean_data(v) for k, v in obj.items()}
            elif isinstance(obj, list):
                return [clean_data(item) for item in obj]
            else:
                return obj

        cleaned_data = clean_data(data)

        # Convert to JSON
        if pretty:
            json_str = json.dumps(cleaned_data, indent=2)
        else:
            json_str = json.dumps(cleaned_data)

        json_bytes = json_str.encode("utf-8")
        logger.info(f"Exported data to JSON ({len(json_bytes)} bytes)")

        return json_bytes

    async def export_to_excel(
        self, data: List[Dict[str, Any]], filename: str = "export.xlsx"
    ) -> bytes:
        """
        Export data to Excel format with styling

        Args:
            data: List of dictionaries to export
            filename: Output filename

        Returns:
            Excel file as bytes
        """
        if not data:
            raise ValueError("No data to export")

        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Dashboard Export"

        # Get headers from first row
        headers = list(data[0].keys())

        # Style for header row
        header_fill = PatternFill(start_color="1976d2", end_color="1976d2", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # Write headers
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header.replace("_", " ").title()
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # Write data rows
        for row_idx, row_data in enumerate(data, start=2):
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = row_data.get(header)

                # Convert datetime to string
                if isinstance(value, datetime):
                    cell.value = value.isoformat()
                elif isinstance(value, dict) or isinstance(value, list):
                    cell.value = json.dumps(value)
                else:
                    cell.value = value

        # Auto-adjust column widths
        for col_idx, header in enumerate(headers, start=1):
            max_length = len(str(header))
            for row_idx in range(2, len(data) + 2):
                cell_value = str(ws.cell(row=row_idx, column=col_idx).value or "")
                max_length = max(max_length, len(cell_value))

            # Set column width (with some padding)
            ws.column_dimensions[chr(64 + col_idx)].width = min(max_length + 2, 50)

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        excel_bytes = output.getvalue()

        logger.info(f"Exported {len(data)} rows to Excel ({len(excel_bytes)} bytes)")

        return excel_bytes

    async def export_dashboard_to_pdf(
        self,
        company_name: str,
        kpi_data: Dict[str, Any],
        revenue_trend: List[Dict[str, Any]],
        recent_activity: List[Dict[str, Any]],
    ) -> bytes:
        """
        Export complete dashboard to PDF report

        Args:
            company_name: Name of the company
            kpi_data: KPI metrics data
            revenue_trend: Revenue trend data
            recent_activity: Recent activity items

        Returns:
            PDF file as bytes
        """
        # Prepare report data for PDF service
        report_data = {
            "summary": f"This dashboard report provides a comprehensive overview of {company_name}'s business performance.",
            "metrics": {
                "Total Revenue": {
                    "value": f"${kpi_data.get('total_revenue', 0):,.2f}",
                    "change": f"{kpi_data.get('revenue_change_percent', 0):+.1f}%",
                },
                "Active Customers": {
                    "value": f"{kpi_data.get('active_customers', 0):,}",
                    "change": f"{kpi_data.get('customers_change_percent', 0):+.1f}%",
                },
                "Inventory Value": {
                    "value": f"${kpi_data.get('inventory_value', 0):,.2f}",
                    "change": f"{kpi_data.get('inventory_change_percent', 0):+.1f}%",
                },
                "Unpaid Invoices": {
                    "value": f"${kpi_data.get('unpaid_invoices', 0):,.2f}",
                    "change": f"{kpi_data.get('invoices_change_percent', 0):+.1f}%",
                },
            },
            "revenue_data": [
                {"date": item["month"], "value": item["revenue"]}
                for item in revenue_trend
            ],
            "transactions": [
                {
                    "date": item.get("timestamp", datetime.now()),
                    "description": item.get("description", "N/A"),
                    "amount": item.get("amount", 0),
                    "status": item.get("status", "N/A"),
                }
                for item in recent_activity[:20]  # Limit to 20 items
            ],
            "insights": [
                f"Data sources: {', '.join(kpi_data.get('data_sources', ['No integrations']))}",
                f"Last updated: {kpi_data.get('last_updated', 'N/A')}",
            ],
        }

        # Generate PDF using PDF service
        pdf_bytes = await pdf_service.generate_business_report(
            company_name=company_name,
            report_type="Dashboard Summary",
            report_data=report_data,
            include_charts=True,
        )

        logger.info(f"Exported dashboard to PDF for {company_name} ({len(pdf_bytes)} bytes)")

        return pdf_bytes

    async def export_transactions_to_csv(
        self, transactions: List[Dict[str, Any]]
    ) -> bytes:
        """
        Export transactions to CSV with proper formatting

        Args:
            transactions: List of transaction dictionaries

        Returns:
            CSV file as bytes
        """
        if not transactions:
            raise ValueError("No transactions to export")

        # Flatten transaction data for CSV
        flattened = []
        for tx in transactions:
            flattened.append({
                "Date": tx.get("timestamp", ""),
                "Type": tx.get("type", ""),
                "Description": tx.get("description", ""),
                "Amount": tx.get("amount", 0),
                "Status": tx.get("status", ""),
                "Source": tx.get("source", ""),
            })

        return await self.export_to_csv(flattened, filename="transactions.csv")

    async def export_analytics_to_excel(
        self, analytics_data: Dict[str, Any]
    ) -> bytes:
        """
        Export analytics data to Excel with multiple sheets

        Args:
            analytics_data: Dictionary of analytics data

        Returns:
            Excel file as bytes
        """
        # Create workbook
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Header styling
        header_fill = PatternFill(start_color="1976d2", end_color="1976d2", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        # Summary sheet
        ws_summary = wb.create_sheet("Summary")
        ws_summary.append(["Metric", "Value"])
        for key, value in analytics_data.get("summary", {}).items():
            ws_summary.append([key.replace("_", " ").title(), value])

        # Style summary headers
        for cell in ws_summary[1]:
            cell.fill = header_fill
            cell.font = header_font

        # Revenue trend sheet (if available)
        if "revenue_trend" in analytics_data:
            ws_revenue = wb.create_sheet("Revenue Trend")
            revenue_data = analytics_data["revenue_trend"]
            if revenue_data:
                headers = list(revenue_data[0].keys())
                ws_revenue.append(headers)
                for row in revenue_data:
                    ws_revenue.append([row.get(h) for h in headers])

                # Style headers
                for cell in ws_revenue[1]:
                    cell.fill = header_fill
                    cell.font = header_font

        # Customers sheet (if available)
        if "customers" in analytics_data:
            ws_customers = wb.create_sheet("Customers")
            customer_data = analytics_data["customers"]
            if customer_data:
                headers = list(customer_data[0].keys())
                ws_customers.append(headers)
                for row in customer_data:
                    ws_customers.append([row.get(h) for h in headers])

                # Style headers
                for cell in ws_customers[1]:
                    cell.fill = header_fill
                    cell.font = header_font

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        logger.info("Exported analytics to multi-sheet Excel file")

        return output.getvalue()


# Singleton instance
export_service = ExportService()
